"""Write and validate paper tables in a 23-worksheet results workbook."""
from __future__ import annotations

from copy import copy
import json
import math
import os
from pathlib import Path
import tempfile
from typing import Any, Mapping
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from .config import ROOT
from .io import atomic_json

SCHEMA = {
    item["name"]: item
    for item in json.loads((ROOT / "prnn/resources/workbook_schema.json").read_text(encoding="utf-8"))
}
SHEET_NAMES = tuple(SCHEMA)
EXPECTED_SHEET_COUNT = 23
if len(SHEET_NAMES) != EXPECTED_SHEET_COUNT:
    raise RuntimeError("The workbook schema must define exactly 23 worksheets.")
ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}
XML_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _value(value: Any) -> Any:
    """Convert array scalars and undefined metrics to Excel-compatible values."""
    if isinstance(value, np.generic):
        value = value.item()
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, str):
        # Some libraries expose version strings as str subclasses. Excel
        # stores ordinary strings; canonicalize before strict read-back checks.
        return str(value) if value else None
    if isinstance(value, float) and not math.isfinite(value):
        # Empty threshold subsets have undefined metrics, not measured zeros.
        return None
    return value


def align_table(name: str, frame: pd.DataFrame) -> pd.DataFrame:
    """Select the paper-table columns in their documented order."""
    if name not in SCHEMA:
        raise ValueError(f"Unknown worksheet: {name}")
    if not isinstance(frame, pd.DataFrame):
        raise TypeError(f"{name} must be a pandas DataFrame.")
    if not frame.columns.is_unique:
        raise ValueError(f"Duplicate columns in {name}.")
    expected = SCHEMA[name]["columns"]
    missing = [column for column in expected if column not in frame.columns]
    if missing:
        raise ValueError(f"Missing columns in {name}: {missing}")
    return frame.loc[:, expected]


def _validate_template(workbook: openpyxl.Workbook) -> None:
    if workbook.sheetnames != list(SHEET_NAMES):
        raise ValueError("Template sheet names/order differ from the required 23-sheet schema.")
    for name in SHEET_NAMES:
        sheet = workbook[name]
        headers = [cell.value for cell in sheet[1]]
        if headers != SCHEMA[name]["columns"]:
            raise ValueError(f"Template headers differ in {name}.")
        if len(sheet.tables) != 1:
            raise ValueError(f"Expected one Excel table in {name}.")
        if sheet.merged_cells.ranges:
            raise ValueError(f"Unexpected merged cells in the one-table worksheet {name}.")


def _styles_for_data(sheet: Any) -> tuple[list[Any], dict[tuple[Any, Any], list[Any]]]:
    defaults = [copy(cell._style) for cell in sheet[2]]
    semantic = {}
    if sheet.title == "Run setup":
        # The Value column contains integers, floats, booleans and text; the
        # cell format follows the parameter rather than its row number.
        for row in sheet.iter_rows(min_row=2):
            semantic[(row[0].value, row[1].value)] = [copy(cell._style) for cell in row]
    return defaults, semantic


def write_workbook(
    tables: Mapping[str, pd.DataFrame],
    output: Path,
    *,
    provenance: str = "PRNN model results",
    source: Path | None = None,
) -> tuple[dict, dict]:
    """Atomically write and read back exactly 23 sheets, one table per sheet.

    Worksheet names/order, A1 headers, table names/styles, column widths,
    freeze panes, header heights and number formats follow the template.
    Row counts and table/filter ranges follow the newly calculated data.
    No navigation sheet, grouped-sheet export, or extra title rows are added.
    """
    if set(tables) != set(SHEET_NAMES):
        raise ValueError(
            "Workbook tables differ from required schema. "
            f"Missing={set(SHEET_NAMES) - set(tables)}; extra={set(tables) - set(SHEET_NAMES)}"
        )
    aligned = {name: align_table(name, tables[name]) for name in SHEET_NAMES}
    for name, frame in aligned.items():
        if len(frame) + 1 > 1_048_576:
            raise ValueError(f"{name} exceeds Excel's worksheet row limit.")
        if frame.empty:
            raise ValueError(f"{name} is empty; refusing to publish an incomplete results workbook.")
    output = Path(output)
    source = Path(source) if source is not None else ROOT / "prnn/resources/workbook_template.xlsx"
    if output.resolve() == source.resolve():
        raise ValueError("Output must not overwrite the workbook template.")
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(source, data_only=False)
    _validate_template(workbook)
    workbook.properties.title = "PRNN model data"
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.subject = provenance
    workbook.properties.description = "23 worksheets; one result table per worksheet. " + provenance
    manifest = {
        "format_version": "23-sheets-v1",
        "sheet_count": EXPECTED_SHEET_COUNT,
        "sheet_names": list(SHEET_NAMES),
        "provenance": provenance,
        "sheets": {},
    }
    for name in SHEET_NAMES:
        sheet = workbook[name]
        frame = aligned[name]
        column_count = len(frame.columns)
        table = next(iter(sheet.tables.values()))
        expected_range = f"A1:{get_column_letter(column_count)}{len(frame) + 1}"
        default_styles, semantic_styles = _styles_for_data(sheet)
        default_height = sheet.sheet_format.defaultRowHeight
        # Clear the template data area before populating results.
        sheet.delete_rows(2, max(sheet.max_row - 1, 0))
        for row_index in list(sheet.row_dimensions):
            if row_index >= 2:
                del sheet.row_dimensions[row_index]
        for row_number, values in enumerate(frame.itertuples(index=False, name=None), 2):
            key = (values[0], values[1]) if name == "Run setup" else None
            styles = semantic_styles.get(key, default_styles)
            for column_number, raw in enumerate(values, 1):
                value = _value(raw)
                cell = sheet.cell(row_number, column_number, value)
                cell._style = copy(styles[column_number - 1])
                if isinstance(value, str):
                    # Text beginning with '=' is a literal, never a formula.
                    cell.data_type = "s"
                if name == "Run setup" and column_number == 3 and key not in semantic_styles:
                    cell.number_format = (
                        "0" if isinstance(value, int) and not isinstance(value, bool)
                        else "0.000000" if isinstance(value, float)
                        else "General"
                    )
            if name in {"Run setup", "Holdout protocols"}:
                # Keep concise setup and protocol descriptions readable.
                lines = 1
                for column_number, raw in enumerate(values, 1):
                    cell = sheet.cell(row_number, column_number)
                    if isinstance(raw, str) and cell.alignment.wrap_text:
                        width = sheet.column_dimensions[get_column_letter(column_number)].width or 13
                        lines = max(lines, sum(max(1, math.ceil(len(part) / max(8, width - 2)))
                                               for part in raw.splitlines()))
                if lines > 1:
                    sheet.row_dimensions[row_number].height = min(409, max(default_height, 13 * lines + 3))
        table.ref = expected_range
        if table.autoFilter is not None:
            table.autoFilter.ref = expected_range
        if table.sortState is not None:
            table.sortState.ref = expected_range
        # Each worksheet has an independent table filter.
        manifest["sheets"][name] = {
            "sheet": name, "header_row": 1, "data_start_row": 2,
            "row_count": len(frame), "column_count": column_count,
            "columns": list(frame.columns), "range": expected_range,
            "table_name": table.displayName,
            "table_style": dict(table.tableStyleInfo) if table.tableStyleInfo else None,
            "freeze_panes": sheet.freeze_panes,
            "header_height": sheet.row_dimensions[1].height,
            "column_widths": {letter: dimension.width for letter, dimension in sheet.column_dimensions.items()},
        }
    descriptor, temp_name = tempfile.mkstemp(prefix=output.stem + ".", suffix=".tmp.xlsx", dir=output.parent)
    os.close(descriptor)
    temporary = Path(temp_name)
    try:
        workbook.save(temporary)
        check = verify_workbook(temporary, manifest, tables=aligned)
        if not check["passed"]:
            raise RuntimeError(f"Workbook verification failed: {check}")
        os.replace(temporary, output)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    internal = output.parent / "_internal"
    atomic_json(internal / (output.stem + ".manifest.json"), manifest)
    atomic_json(internal / (output.stem + ".verification.json"), check)
    return manifest, check


def verify_workbook(path: Path, manifest: dict, tables: Mapping[str, pd.DataFrame] | None = None) -> dict:
    """Check saved values, dimensions, headers and all 23 Excel table ranges."""
    errors = []
    error_count = 0
    compared = nonempty = checked_tables = formula_count = 0

    def fail(message: str) -> None:
        nonlocal error_count
        error_count += 1
        if len(errors) < 20:
            errors.append(message)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    actual_sheets = list(workbook.sheetnames)
    if actual_sheets != list(SHEET_NAMES):
        fail(f"Expected exactly 23 ordered sheets; found {actual_sheets}")
    if list(manifest.get("sheets", {})) != list(SHEET_NAMES):
        fail("The manifest does not contain the required 23 ordered sheets.")
    for name in SHEET_NAMES:
        if name not in workbook.sheetnames or name not in manifest["sheets"]:
            fail(f"Missing worksheet: {name}")
            continue
        spec = manifest["sheets"][name]
        sheet = workbook[name]
        checked_tables += 1
        if sheet.max_row != spec["row_count"] + 1 or sheet.max_column != spec["column_count"]:
            fail(f"Unexpected dimensions or stale trailing data in {name}.")
        expected_rows = (align_table(name, tables[name]).itertuples(index=False, name=None)
                         if tables is not None else None)
        for row_number, cells in enumerate(sheet.iter_rows(), 1):
            if row_number == 1:
                if [cell.value for cell in cells] != spec["columns"]:
                    fail(f"Column header/order mismatch in {name}.")
                continue
            truth = next(expected_rows, None) if expected_rows is not None else None
            if expected_rows is not None and truth is None:
                fail(f"Extra row {row_number} in {name}.")
            for column_number, cell in enumerate(cells, 1):
                actual = cell.value
                if actual is not None:
                    nonempty += 1
                if cell.data_type == "f":
                    formula_count += 1
                    fail(f"Unexpected formula in exported result: {name}!{cell.coordinate}")
                if cell.data_type == "e" or (isinstance(actual, str) and actual in ERRORS):
                    fail(f"Excel error in {name}!{cell.coordinate}: {actual}")
                if truth is None:
                    continue
                expected = _value(truth[column_number - 1])
                compared += 1
                if (isinstance(actual, (int, float)) and not isinstance(actual, bool)
                        and isinstance(expected, (int, float)) and not isinstance(expected, bool)):
                    equal = math.isclose(actual, expected, rel_tol=2e-14, abs_tol=1e-15)
                else:
                    equal = type(actual) is type(expected) and actual == expected
                if not equal:
                    fail(f"Value mismatch in {name}!{cell.coordinate}: {actual!r} != {expected!r}")
    workbook.close()
    # Table/filter XML is small; inspect it without loading 900,000 styled cells.
    with ZipFile(path) as archive:
        table_paths = [name for name in archive.namelist() if name.startswith("xl/tables/") and name.endswith(".xml")]
        saved_tables = {}
        for table_path in table_paths:
            element = ET.fromstring(archive.read(table_path))
            saved_tables[element.attrib["displayName"]] = element
        if len(saved_tables) != EXPECTED_SHEET_COUNT:
            fail(f"Expected 23 Excel tables, found {len(saved_tables)}.")
        for name, spec in manifest["sheets"].items():
            element = saved_tables.get(spec["table_name"])
            if element is None:
                fail(f"Missing Excel table in {name}.")
                continue
            if element.attrib.get("ref") != spec["range"]:
                fail(f"Table range mismatch in {name}.")
            columns = element.find("m:tableColumns", XML_NS)
            if columns is None or [c.attrib["name"] for c in columns] != spec["columns"]:
                fail(f"Excel table column definitions differ in {name}.")
            auto_filter = element.find("m:autoFilter", XML_NS)
            if auto_filter is not None and auto_filter.attrib.get("ref") != spec["range"]:
                fail(f"Filter range mismatch in {name}.")
    if tables is not None:
        expected_count = sum(len(tables[name]) * len(SCHEMA[name]["columns"]) for name in SHEET_NAMES)
        if compared != expected_count:
            fail(f"Compared {compared} data cells instead of {expected_count}.")
    return {
        "passed": error_count == 0, "sheet_count": len(actual_sheets),
        "sheets": actual_sheets, "table_count": checked_tables,
        "cells_compared": compared, "nonempty_data_cells": nonempty,
        "unexpected_formulas": formula_count, "error_count": error_count,
        "numeric_relative_tolerance": 2e-14, "numeric_absolute_tolerance": 1e-15,
        "errors": errors,
    }


def writer_preflight(directory: Path) -> None:
    """Fail before training when openpyxl or the output directory is unusable."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    handle, path = tempfile.mkstemp(suffix=".xlsx", dir=directory)
    os.close(handle)
    try:
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "writer_preflight"
        workbook.save(path)
        workbook.close()
        check = openpyxl.load_workbook(path, read_only=True)
        okay = check.active["A1"].value == "writer_preflight"
        check.close()
        if not okay:
            raise RuntimeError("Excel writer preflight failed.")
    finally:
        Path(path).unlink(missing_ok=True)
